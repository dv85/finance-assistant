"""All fixtures are fabricated in memory or in a temporary workspace."""
from __future__ import annotations

import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from household_config import ROOT, load_config
from import_formats import csv_rows, load_mapping, xlsx_rows
from convert_bank_exports import normalize_row, prepare_import
from ledger_memory import digest, empty_ledger, to_app_state, validate_ledger
from build_sites_import import build_transactions
from sync_local_ledger import LocalClient, publish, validate_url


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.config = load_config(ROOT / "examples/household.example.json")
        self.mapping = load_mapping(ROOT / "examples/csv-mapping.example.json")

    def parse(self, text, mapping=None):
        path = self.folder / "synthetic.csv"
        path.write_text(text, encoding="utf-8")
        return csv_rows(path, mapping or self.mapping)

    def raw(self, **changes):
        return {"date": "2025-04-02", "description": "Example Market", "signedAmount": -10, "reference": "example-1",
                "source": "synthetic.csv", "sourceRow": 2, **changes}

    def test_basic_signed_csv(self):
        rows, audit = self.parse("Date,Description,Amount,Type,Category,Reference\n2025-04-02,Example Market,-10.50,expense,Groceries,synthetic-1\n")
        self.assertEqual(rows[0]["signedAmount"], -10.5)
        self.assertEqual(audit["errors"], [])

    def test_split_debit_credit_and_localized_decimals(self):
        mapping = {**self.mapping, "delimiter": ";", "decimalSeparator": ",", "thousandsSeparator": ".",
                   "dateFormat": "%d/%m/%Y", "columns": {"date": "Date", "description": "Description", "debit": "Debit", "credit": "Credit"}}
        rows, audit = self.parse("Date;Description;Debit;Credit\n02/04/2025;Example Market;1.250,50;\n03/04/2025;Example Payroll;;2.000,00\n", mapping)
        self.assertEqual([row["signedAmount"] for row in rows], [-1250.5, 2000])
        self.assertFalse(audit["errors"])

    def test_debit_positive(self):
        rows, _ = self.parse("Date,Description,Amount,Type,Category,Reference\n2025-04-02,Example Market,10,expense,Groceries,x\n", {**self.mapping, "signConvention": "debit_positive"})
        self.assertEqual(rows[0]["signedAmount"], -10)

    def test_invalid_dates_amounts_and_malformed_rows_fail_preview(self):
        for value in ["2025-02-30,Example Market,-10,expense,Groceries,x", "2025-04-02,Example Market,NaN,expense,Groceries,x",
                      "2025-04-02,Example Market,-10.123,expense,Groceries,x", "2025-04-02,Example Market,-10,expense,Groceries,x,extra"]:
            with self.subTest(value=value):
                rows, audit = self.parse("Date,Description,Amount,Type,Category,Reference\n" + value + "\n")
                _, report = prepare_import(empty_ledger(self.config), self.config, rows, audit, "checking", "synthetic-hash")
                self.assertEqual(report["status"], "fail")

    def test_duplicate_headers_are_rejected(self):
        with self.assertRaises(ValueError):
            self.parse("Date,Description,Amount,Type,Category,Reference,Date\n")

    def test_duplicate_import_is_idempotent(self):
        raw = self.raw()
        ledger, first = prepare_import(empty_ledger(self.config), self.config, [raw], {"errors": []}, "checking", "a")
        second, report = prepare_import(ledger, self.config, [raw], {"errors": []}, "checking", "a")
        self.assertEqual((first["inserted"], report["inserted"], report["duplicates"]), (1, 0, 1))
        self.assertEqual(ledger, second)

    def test_identical_unreferenced_purchases_are_preserved(self):
        raw = self.raw(reference="")
        ledger, report = prepare_import(empty_ledger(self.config), self.config, [raw, {**raw, "sourceRow": 3}], {"errors": []}, "checking", "a")
        self.assertEqual(len(ledger["transactions"]), 2)
        again, _ = prepare_import(ledger, self.config, [raw, {**raw, "sourceRow": 3}], {"errors": []}, "checking", "a")
        self.assertEqual(len(again["transactions"]), 2)

    def test_refund_income_transfer_and_settlement_are_distinct(self):
        rows = [self.raw(reference="refund", signedAmount=10, type="refund"), self.raw(reference="income", signedAmount=100, type="income"),
                self.raw(reference="transfer", type="transfer"), self.raw(reference="settlement", type="card_settlement")]
        ledger, report = prepare_import(empty_ledger(self.config), self.config, rows, {"errors": []}, "checking", "a")
        self.assertEqual(report["status"], "pass")
        self.assertEqual({row["type"] for row in ledger["transactions"]}, {"refund", "income", "transfer", "card_settlement"})
        self.assertEqual(len(to_app_state(ledger, self.config, "2025-04")["income"]), 1)

    def test_currency_unknown_account_and_direction_mismatch_fail(self):
        for row in [self.raw(currency="EUR"), self.raw(account="not-configured"), self.raw(type="income"), self.raw(signedAmount=0)]:
            _, report = prepare_import(empty_ledger(self.config), self.config, [row], {"errors": []}, "checking", "a")
            self.assertEqual(report["status"], "fail")

    def test_installments_are_explicit_and_roundtrip(self):
        first = self.raw(account="card", isInstallment=True, purchaseAmount=60, installmentNumber=1, installmentCount=2, signedAmount=-30,
                         date="2025-04-02", billingDate="2025-05-10", reference="demo-plan")
        second = {**first, "sourceRow": 3, "billingDate": "2025-06-10", "installmentNumber": 2}
        ledger, report = prepare_import(empty_ledger(self.config), self.config, [first, second], {"errors": []}, None, "a")
        self.assertEqual(report["status"], "pass")
        self.assertEqual(len(ledger["installmentPlans"]), 1)
        self.assertEqual(sum(item["amount"] for item in ledger["installmentPlans"][0]["schedule"]), 60)
        wire = build_transactions(ledger)
        self.assertEqual(wire[0]["reportingDate"], wire[0]["billingDate"])
        self.assertEqual(wire[0]["transactionDate"], "2025-04-02")

    def test_installment_requires_billing_date_and_valid_ordinals(self):
        for changes in [{"isInstallment": True}, {"isInstallment": True, "billingDate": "2025-05-10", "installmentNumber": 3, "installmentCount": 2}]:
            with self.assertRaises(ValueError):
                normalize_row(self.raw(**changes), self.config, "card", 1)

    def test_repeated_amounts_do_not_infer_installments(self):
        ledger, report = prepare_import(empty_ledger(self.config), self.config, [self.raw(reference="a"), self.raw(reference="b")], {"errors": []}, "card", "a")
        self.assertFalse(ledger["installmentPlans"])
        self.assertEqual(report["status"], "pass")

    def test_missing_installment_count_is_not_inferred_from_principal(self):
        row = self.raw(account="card", isInstallment=True, purchaseAmount=120, signedAmount=-30,
                       billingDate="2025-05-10", creditTransactionType="תשלומים רגיל")
        ledger, report = prepare_import(empty_ledger(self.config), self.config, [row], {"errors": []}, None, "a")
        self.assertEqual(report["status"], "pass")
        self.assertEqual(ledger["installmentPlans"][0]["status"], "needs_review")
        self.assertIsNone(ledger["installmentPlans"][0]["installmentCount"])

    def test_preview_does_not_mutate_input_ledger(self):
        ledger = empty_ledger(self.config)
        before = copy.deepcopy(ledger)
        prepare_import(ledger, self.config, [self.raw()], {"errors": []}, "checking", "a")
        self.assertEqual(ledger, before)

    def test_israeli_bank_xlsx_balances(self):
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["Synthetic statement"])
        for _ in range(4): sheet.append([])
        sheet.append(["01/04/2025", "Example Payroll", "", "bank-a", 0, 100, 1100, "", "", ""])
        sheet.append(["02/04/2025", "Example Market", "", "bank-b", 10, 0, 1090, "", "", ""])
        path = self.folder / "synthetic-bank.xlsx"; workbook.save(path)
        config = {**self.config, "currency": "ILS"}
        rows, audit = xlsx_rows(path, "israeli-bank", config, "oldest-first")
        self.assertEqual([row["signedAmount"] for row in rows], [100, -10])
        self.assertFalse(audit["errors"])
        sheet.cell(7, 7, 1000); workbook.save(path)
        with self.assertRaisesRegex(ValueError, "reconcile"):
            xlsx_rows(path, "israeli-bank", config, "oldest-first")

    def test_israeli_card_xlsx(self):
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["פירוט עבור כרטיס"])
        sheet.append(["שם כרטיס", "חיוב לתאריך", "תאריך", "שם בית עסק", 'סכום חיוב בש"ח', "סכום קנייה", "מטבע", "אסמכתא", "תאור סוג עסקת אשראי"])
        sheet.append(["9999", "10/05/2025", "02/04/2025", "Example Market", 30, 30, "ILS", "synthetic-a", "רגיל"])
        path = self.folder / "synthetic-card.xlsx"; workbook.save(path)
        rows, _ = xlsx_rows(path, "israeli-card", {**self.config, "currency": "ILS"})
        self.assertEqual(rows[0]["signedAmount"], -30)
        self.assertFalse(rows[0]["isInstallment"])


class CliTests(unittest.TestCase):
    def test_preview_apply_stale_guard_and_setup_preservation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            shutil.copytree(ROOT / "scripts", root / "scripts", ignore=shutil.ignore_patterns("__pycache__"))
            shutil.copytree(ROOT / "examples", root / "examples")
            def run(*arguments):
                return subprocess.run([sys.executable, *arguments], cwd=root, capture_output=True, text=True)
            self.assertEqual(run("scripts/setup_local.py").returncode, 0)
            ledger_path = root / "memory/ledger.json"
            before = ledger_path.read_bytes()
            command = ["scripts/convert_bank_exports.py", "--format", "csv", "--input", "examples/transactions.synthetic.csv", "--mapping", "examples/csv-mapping.example.json", "--account", "checking", "--month", "2025-04"]
            self.assertNotEqual(run(*command, "--apply").returncode, 0)
            self.assertEqual(run(*command).returncode, 0)
            self.assertEqual(ledger_path.read_bytes(), before)
            config = root / "config/household.json"
            original = config.read_bytes(); config.write_bytes(original + b"\n")
            self.assertNotEqual(run(*command, "--apply").returncode, 0)
            config.write_bytes(original)
            applied = run(*command, "--apply")
            self.assertEqual(applied.returncode, 0, applied.stderr)
            ledger = json.loads(ledger_path.read_text()); validate_ledger(ledger)
            self.assertEqual(len(ledger["transactions"]), 5)
            snapshot = ledger_path.read_bytes()
            self.assertEqual(run("scripts/setup_local.py").returncode, 0)
            self.assertEqual(ledger_path.read_bytes(), snapshot)
            self.assertNotEqual(run("scripts/setup_local.py", "--demo").returncode, 0)


class SyncTests(unittest.TestCase):
    def test_loopback_only(self):
        for value in ["http://127.0.0.1:4310", "http://localhost:4310", "http://[::1]:4310"]:
            self.assertEqual(validate_url(value), value)
        for value in ["https://example.com:443", "http://127.0.0.1:4310/path", "http://localhost.evil.invalid:4310", "http://user@localhost:4310", "http://localhost", "http://localhost:4310?x=y"]:
            with self.assertRaises(ValueError): validate_url(value)

    def test_staging_failure_rolls_back_without_commit(self):
        calls = []
        class Fake:
            def call(self, path, method="GET", body=None):
                calls.append((path, method))
                if path.endswith("/start"): return {"snapshotId": "synthetic"}
                if path.endswith("/chunk"): raise ValueError("synthetic chunk rejection")
                return {}
        bundle = {"schemaVersion": "finance-sites-import/v1", "sourceHash": "synthetic", "createdAt": "2025-04-01", "state": {}, "expected": {},
                  "accounts": [{"id": "example"}], "transactions": [], "balanceSnapshots": [], "importRuns": []}
        with self.assertRaises(ValueError): publish(Fake(), bundle)
        self.assertIn(("/api/admin/imports/synthetic", "DELETE"), calls)
        self.assertFalse(any(path.endswith("/commit") for path, _ in calls))

    def test_active_verification_failure_is_reported(self):
        class Fake:
            def call(self, path, method="GET", body=None):
                return {"snapshotId": "synthetic"} if path.endswith("/start") else {}
        bundle = {"schemaVersion": "finance-sites-import/v1", "sourceHash": "synthetic", "createdAt": "2025-04-01", "state": {}, "expected": {},
                  "accounts": [], "transactions": [], "balanceSnapshots": [], "importRuns": []}
        with self.assertRaisesRegex(ValueError, "verification"):
            publish(Fake(), bundle)


class PublicationTests(unittest.TestCase):
    def test_fresh_database_migration_contains_schema_only(self):
        import sqlite3
        connection = sqlite3.connect(":memory:")
        self.addCleanup(connection.close)
        connection.executescript((ROOT / "web/drizzle/0001_initial.sql").read_text())
        tables = [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
        self.assertGreater(len(tables), 10)
        for table in tables:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM "' + table + '"').fetchone()[0], 0)

    def test_manifest_refuses_private_traversal_missing_files_and_symlinks(self):
        from public_release import public_files
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            source = root / "public.txt"; source.write_text("synthetic text")
            manifest = root / "PUBLIC_FILES.json"
            for names in [["../outside.txt"], ["memory/ledger.json"], [".env"], ["missing.txt"], ["public.txt", "public.txt"]]:
                manifest.write_text(json.dumps(names))
                with self.assertRaises(ValueError): public_files(root)
            manifest.write_text(json.dumps(["public.txt"]))
            self.assertEqual(public_files(root), [source])
            alias = root / "alias.txt"; alias.symlink_to(source)
            manifest.write_text(json.dumps(["alias.txt"]))
            with self.assertRaises(ValueError): public_files(root)

    def test_the_exact_public_set_has_no_builtin_privacy_findings(self):
        from public_release import audit
        files, findings = audit()
        self.assertGreater(len(files), 50)
        self.assertEqual(findings, [])


if __name__ == "__main__":
    unittest.main()
