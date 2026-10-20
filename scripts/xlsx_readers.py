"""Format readers extracted from the existing XLSX importer; no household rules or FX estimates."""
from __future__ import annotations
import hashlib
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

CURRENCY = "ILS"
PARSE_AUDIT: dict = {}
CARD_ADJUSTMENTS: list = []
PENDING_CARD_FX_RATES: dict = {}

def diners_statement_fx_rates(_worksheet):
    # Unbilled foreign amounts require review rather than an invented conversion.
    return {}

def file_checksum(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def audit_for(path: Path, parser: str) -> dict[str, Any]:
    key = str(path.resolve())
    if key not in PARSE_AUDIT:
        PARSE_AUDIT[key] = {
            "file": path.name,
            "path": key,
            "checksum": file_checksum(path),
            "parser": parser,
            "dateRange": {"from": None, "to": None},
            "rawRows": 0,
            "acceptedRows": 0,
            "duplicateRows": 0,
            "skippedRowsByReason": {},
            "invalidRows": 0,
            "amountTotals": {"accepted": 0.0, "duplicates": 0.0, "skipped": 0.0, "invalid": 0.0},
        }
    return PARSE_AUDIT[key]


def count_skip(audit: dict[str, Any], reason: str, amount: float = 0.0) -> None:
    audit["skippedRowsByReason"][reason] = audit["skippedRowsByReason"].get(reason, 0) + 1
    audit["amountTotals"]["skipped"] = round(audit["amountTotals"]["skipped"] + abs(amount), 2)


def observe_date(audit: dict[str, Any], value: str) -> None:
    if not value:
        return
    current = audit["dateRange"]
    current["from"] = value if current["from"] is None else min(current["from"], value)
    current["to"] = value if current["to"] is None else max(current["to"], value)


def workbook(path: Path):
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "Missing dependency: openpyxl. Run scripts/setup_local.py "
            "or install openpyxl in your active Python environment."
        ) from exc

    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def iso(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    if not text:
        return ""

    iso_match = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", text)
    if iso_match:
        return "-".join(iso_match.groups())

    local_match = re.search(r"\b(\d{1,2})[./](\d{1,2})[./](\d{4})\b", text)
    if local_match:
        day, month, year = local_match.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"

    return ""


def num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    return round(float(value), 2)


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def ref_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def uid(prefix: str, *parts: Any) -> str:
    raw = "|".join(clean(part) for part in parts)
    return f"{prefix}-{hashlib.sha1(raw.encode('utf-8')).hexdigest()[:14]}"


def parse_bank_transactions(path: Path) -> list[dict[str, Any]]:
    source_path = path
    audit = audit_for(source_path, "israeli-bank-xlsx-v2")
    wb = workbook(source_path)
    ws = wb.worksheets[0]
    rows = []
    for row_index, row in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        row = tuple(row) + (None,) * max(0, 10 - len(row))
        tx_date = iso(row[0])
        if not tx_date:
            continue
        audit["rawRows"] += 1
        observe_date(audit, tx_date)
        action = clean(row[1])
        details = clean(row[2])
        reference = ref_text(row[3])
        debit = num(row[4])
        credit = num(row[5])
        balance = num(row[6])
        beneficiary = clean(row[8])
        purpose = clean(row[9])
        if not action or (debit == 0 and credit == 0):
            count_skip(audit, "non_movement_row", debit or credit)
            continue
        rows.append({
            "row": row_index,
            "source_file": source_path.name,
            "date": tx_date,
            "action": action,
            "details": details,
            "reference": reference,
            "debit": debit,
            "credit": credit,
            "balance": balance,
            "beneficiary": beneficiary,
            "purpose": purpose,
        })
    return rows


def header_index(header: list[str], *needles: str, exact: bool = False) -> int | None:
    for index, value in enumerate(header):
        if exact and value in needles:
            return index
        if not exact and any(needle in value for needle in needles):
            return index
    return None


def cell(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def card_from_diners_title(text: str) -> str:
    ending_match = re.search(r"(?:המסתיים|מסתיים)\s*ב?-?(\d{4})", text)
    if ending_match:
        return ending_match.group(1)
    match = re.search(r"(\d{4})", text)
    return match.group(1) if match else "unknown"


def billing_date_from_diners_text(text: str) -> str:
    match = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text)
    if not match:
        return ""
    day, month, year = match.groups()
    return f"{year}-{int(month):02d}-{int(day):02d}"


def pending_currency_from_diners_text(text: str) -> str:
    if "בתהליך קליטה" not in text:
        return ""
    match = re.search(r"בנוסף\s+[\d,.]+\s+([A-Z]{3})", text)
    return match.group(1) if match else ""


def normalized_currency(value: Any) -> str:
    currency = clean(value).upper()
    return "RSD" if currency == "RDS" else currency


def currency_from_number_format(number_format: Any) -> str:
    """Read the purchase currency encoded in a Diners workbook cell format."""
    value = clean(number_format).upper()
    if "RSD" in value or "RDS" in value:
        return "RSD"
    if "€" in value or "EUR" in value:
        return "EUR"
    if "₪" in value or "ILS" in value:
        return CURRENCY
    if "GEL" in value:
        return "GEL"
    if "$" in value or "USD" in value:
        return "USD"
    return ""


def parse_diners_card_transactions(path: Path, ws: Any) -> list[dict[str, Any]]:
    rows = []
    audit = audit_for(path, "diners-card-xlsx-v2")
    title = clean(ws.cell(row=1, column=2).value) or clean(ws.cell(row=1, column=1).value)
    card = card_from_diners_title(title)
    billing_date = ""
    pending_currency = ""
    statement_fx_rates = diners_statement_fx_rates(ws)
    header: list[str] | None = None

    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        first = clean(row[0])
        if "עסקאות לחיוב" in first:
            billing_date = billing_date_from_diners_text(first) or billing_date
            continue
        if "עסקאות בתהליך קליטה" in first:
            pending_currency = pending_currency_from_diners_text(first) or pending_currency
            continue
        row_values = [clean(item) for item in row]
        if any("תאריך" in item and "עסקה" in item for item in row_values) and any("שם בית עסק" in item for item in row_values):
            header = row_values
            continue
        if not header:
            continue

        date_index = header_index(header, "תאריך עסקה", "תאריך", exact=True)
        merchant_index = header_index(header, "שם בית עסק")
        purchase_amount_index = header_index(header, "סכום עסקה", "סכום קנייה")
        charge_ils_index = header_index(header, "סכום חיוב")
        type_index = header_index(header, "סוג עסקה")
        branch_index = header_index(header, "ענף")
        note_index = header_index(header, "הערות")

        purchase_date = iso(cell(row, date_index))
        merchant = clean(cell(row, merchant_index))
        purchase_amount = num(cell(row, purchase_amount_index))
        charged_ils = num(cell(row, charge_ils_index))
        transaction_type = clean(cell(row, type_index))
        branch = clean(cell(row, branch_index))
        note = clean(cell(row, note_index))

        if not purchase_date or not merchant:
            continue
        audit["rawRows"] += 1
        observe_date(audit, purchase_date)
        if "שירות אשראי מתגלגל" in merchant:
            count_skip(audit, "non_spending_card_adjustment", charged_ils or purchase_amount)
            CARD_ADJUSTMENTS.append({
                "card": card, "billingDate": billing_date, "amount": charged_ils or purchase_amount,
                "description": merchant, "source": f"{path.name}: row {row_index}",
            })
            continue
        purchase_cell = ws.cell(row=row_index, column=(purchase_amount_index or 0) + 1)
        original_currency = normalized_currency(currency_from_number_format(purchase_cell.number_format)) or CURRENCY
        is_pending = note == "עסקה בקליטה"
        amount_estimated = False
        estimated_fx_rate = None
        estimation_method = None
        if charged_ils == 0:
            if original_currency == CURRENCY and is_pending and pending_currency:
                original_currency = normalized_currency(pending_currency)
            fx_rate = statement_fx_rates.get(original_currency) or PENDING_CARD_FX_RATES.get(original_currency)
            if original_currency != CURRENCY and not fx_rate:
                raise ValueError(
                    f"No settled base-currency amount for {original_currency} row {row_index} in {path.name}"
                )
            charged_ils = round(purchase_amount * fx_rate, 2) if fx_rate else purchase_amount
            amount_estimated = bool(is_pending and original_currency != CURRENCY)
            if amount_estimated:
                estimated_fx_rate = round(float(fx_rate), 8)
                estimation_method = (
                    "same_statement_settled_median"
                    if original_currency in statement_fx_rates
                    else "configured_pending_fx_rate"
                )
        if charged_ils == 0:
            count_skip(audit, "zero_amount", purchase_amount)
            continue

        rows.append({
            "row": row_index,
            "date": purchase_date,
            "billing_date": billing_date,
            "card": card,
            "merchant": merchant,
            "amount": charged_ils,
            "purchase_amount": purchase_amount or charged_ils,
            "original_currency": original_currency,
            "reference": note or f"diners-row-{row_index}",
            "section": title or "Diners statement",
            "source_file": path.name,
            "provider_category": branch,
            "provider_type": transaction_type,
            "is_pending": is_pending,
            "amount_estimated": amount_estimated,
            "estimated_fx_rate": estimated_fx_rate,
            "estimation_method": estimation_method,
        })
    return rows


def parse_card_transactions(path: Path) -> list[dict[str, Any]]:
    wb = workbook(path)
    ws = wb.worksheets[0]
    rows = []
    audit = audit_for(path, "card-xlsx-v2")
    section = ""
    header: list[str] | None = None
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        first = clean(row[0])
        if first.startswith("פירוט עבור"):
            section = first
            header = None
            continue
        if section and first == "שם כרטיס":
            header = [clean(cell) for cell in row]
            continue
        if not section or not header or not first.isdigit():
            continue

        card_index = header_index(header, "שם כרטיס", exact=True)
        billing_index = header_index(header, "חיוב לתאריך")
        date_index = header_index(header, "תאריך", exact=True)
        merchant_index = header_index(header, "שם בית עסק")
        charge_ils_index = header_index(header, "סכום חיוב בש", "סכום חיוב")
        purchase_amount_index = header_index(header, "סכום קנייה")
        currency_index = header_index(header, "מטבע מקורי", "מטבע")
        reference_index = header_index(header, "אסמכתא")
        credit_type_index = header_index(header, "תאור סוג עסקת אשראי", "תיאור סוג עסקת אשראי")

        card = ref_text(cell(row, card_index) or first)
        billing_date = iso(cell(row, billing_index))
        purchase_date = iso(cell(row, date_index))
        if purchase_date.startswith("1900") or not purchase_date:
            purchase_date = billing_date
        merchant = clean(cell(row, merchant_index))
        audit["rawRows"] += 1
        observe_date(audit, purchase_date or billing_date)
        if "שירות אשראי מתגלגל" in merchant:
            count_skip(audit, "non_spending_card_adjustment", num(cell(row, charge_ils_index)))
            CARD_ADJUSTMENTS.append({
                "card": card, "billingDate": billing_date, "amount": num(cell(row, charge_ils_index)),
                "description": merchant, "source": f"{path.name}: row {row_index}",
            })
            continue
        charged_ils = num(cell(row, charge_ils_index))
        purchase_amount = num(cell(row, purchase_amount_index)) or charged_ils
        original_currency = clean(cell(row, currency_index)) or CURRENCY
        reference = ref_text(cell(row, reference_index))
        is_pending = reference in {"", "0"} and "~" in merchant

        if charged_ils == 0:
            count_skip(audit, "zero_amount", purchase_amount)
            continue
        rows.append({
            "row": row_index,
            "date": purchase_date,
            "billing_date": billing_date,
            "card": card,
            "merchant": merchant or "Card charge",
            "amount": charged_ils,
            "purchase_amount": purchase_amount,
            "original_currency": original_currency,
            "reference": reference,
            "credit_transaction_type": clean(cell(row, credit_type_index)),
            "section": section,
            "source_file": path.name,
            "is_pending": is_pending,
            "amount_estimated": False,
            "estimated_fx_rate": None,
            "estimation_method": None,
        })
    if rows:
        return rows
    # The Diners parser owns the accounting for files in its statement format.
    PARSE_AUDIT.pop(str(path.resolve()), None)
    return parse_diners_card_transactions(path, ws)


